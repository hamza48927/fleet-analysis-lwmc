from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).parent
ASSET_DIR = ROOT / "manual_assets"
ASSET_DIR.mkdir(exist_ok=True)
MANUAL = ROOT / "LWMC_Fleet_Dashboard_User_Manual.docx"
QUICK = ROOT / "LWMC_Fleet_Dashboard_Quick_Reference.docx"

NAVY = "0D1F2D"
INK = "1B2B34"
MUTED = "5C6F78"
GREEN = "158B5D"
GREEN_DARK = "0C5D43"
MINT = "E1F3EB"
AMBER = "E2972F"
RED = "B93A39"
LINE = "DAE4E0"
PAPER = "F8FAF8"


def font(size=18, bold=False, color=INK, name="Aptos"):
    f = FontSpec = None
    return None


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def setup_document(doc, title):
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 34, NAVY), ("Heading 1", 23, NAVY),
                              ("Heading 2", 15, GREEN_DARK), ("Heading 3", 11.5, GREEN)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(13 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
    header = section.header.paragraphs[0]
    header.text = "LWMC FLEET DASHBOARD  |  OPERATOR MANUAL"
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.bold = True
    header.runs[0].font.color.rgb = RGBColor.from_string(GREEN_DARK)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    doc.core_properties.title = title
    doc.core_properties.subject = "LWMC Fleet Dashboard operating instructions"
    doc.core_properties.author = "LWMC Fleet Dashboard"


def add_title_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("LWMC OPERATIONS")
    r.font.name = "Aptos"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(GREEN)
    p = doc.add_paragraph(style="Title")
    p.add_run("Fleet Dashboard")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run("Complete user manual")
    r.font.size = Pt(20)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("Live fleet visibility, vehicle-level investigation, and operational reporting for Lahore.")
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(INK)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(80)
    r = p.add_run("Version 1.0  |  11 September 2026")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(MUTED)
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, MINT)
    set_cell_border(cell, MINT)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell.text = "Scope: Lahore city fleet operations. This manual describes the packaged desktop application and its daily data workflow."
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string(GREEN_DARK)
    cell.paragraphs[0].runs[0].font.bold = True
    doc.add_page_break()


def add_callout(doc, label, text, color=GREEN):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, MINT if color == GREEN else "FFF4DF")
    set_cell_border(cell, MINT if color == GREEN else "F2D6A0")
    p = cell.paragraphs[0]
    r = p.add_run(label.upper() + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, NAVY)
        cell.text = head
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            if ridx % 2 == 0:
                set_cell_shading(cells[i], "F1F6F3")
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_figure(doc, path, caption, width=6.7):
    if path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        cap = doc.add_paragraph(caption)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].italic = True
        cap.runs[0].font.size = Pt(8.5)
        cap.runs[0].font.color.rgb = RGBColor.from_string(MUTED)


def draw_ui_screenshot(path, mode):
    width, height = 1500, 850
    image = Image.new("RGB", (width, height), "#f8faf8")
    draw = ImageDraw.Draw(image)
    try:
        bold = ImageFont.truetype("segoeuib.ttf", 25)
        regular = ImageFont.truetype("segoeui.ttf", 17)
        small = ImageFont.truetype("segoeui.ttf", 14)
    except OSError:
        bold = regular = small = ImageFont.load_default()
    draw.rectangle((0, 0, width, 72), fill="#0d1f2d")
    draw.text((28, 18), "LWMC Fleet Dashboard", fill="white", font=bold)
    for x, label in ((420, "Map & Analytics"), (610, "Fleet Table"), (770, "Reports")):
        draw.text((x, 25), label, fill="#b9d5c8" if label != mode else "#7ce0ae", font=regular)
    draw.rectangle((0, 72, width, 122), fill="#f3f7f5", outline="#d7e3de")
    draw.text((25, 88), "Data sources:", fill="#33403c", font=regular)
    labels = ["VTMS Export", "Vehicle Status", "Employee Roster", "Trip Reports", "Fleet Registry"]
    x = 170
    for label in labels:
        draw.rounded_rectangle((x, 84, x + 170, 111), radius=5, fill="#ffffff", outline="#b7c9c2")
        draw.text((x + 10, 91), label, fill="#33403c", font=small)
        x += 190
    draw.rectangle((0, 122, 300, height), fill="#eef6f2")
    if mode == "Fleet Table":
        draw.text((24, 150), "Fleet Table Filters", fill="#0d1f2d", font=bold)
        fields = ["Status: All", "Vehicle Type: All", "Town: All", "Zone: All", "UC: All", "Employee: All staff / vehicles"]
        y = 205
        for label in fields:
            draw.rounded_rectangle((22, y, 278, y + 48), radius=5, fill="white", outline="#c9d8d2")
            draw.text((35, y + 13), label, fill="#33403c", font=small)
            y += 62
        draw.text((24, 610), "3,047 vehicles shown", fill="#0c5d43", font=regular)
        draw.text((24, 645), "Live now: 3,047  (Moving 1,428 - Still 1,619)", fill="#557b70", font=small)
        draw.text((24, 678), "In registry: 3,047  - Registry-only: 0", fill="#557b70", font=small)
        draw.rectangle((300, 122, width, 174), fill="#ffffff", outline="#d7e3de")
        draw.text((325, 145), "Vehicle", fill="#557b70", font=small)
        draw.text((590, 145), "Town", fill="#557b70", font=small)
        draw.text((825, 145), "Type", fill="#557b70", font=small)
        draw.text((1050, 145), "Status", fill="#557b70", font=small)
        draw.text((1240, 145), "Work (min)", fill="#557b70", font=small)
        rows = [("CA-161", "Allama Iqbal Town", "Chain Arm Roll", "moving", "248"), ("C-1390", "Nishtar Town", "Compactor (13CM)", "still", "438"), ("HA-32", "DGBT", "Arm Roll Truck", "still", "200"), ("CA-159", "DGBT", "-", "moving", "309")]
        y = 194
        for i, row in enumerate(rows):
            if i % 2 == 0:
                draw.rectangle((300, y - 4, width, y + 42), fill="#f7faf8")
            for px, val in zip((325, 590, 825, 1050, 1240), row):
                color = "#16a34a" if val == "moving" else "#d97706" if val == "still" else "#33403c"
                draw.text((px, y + 8), val, fill=color, font=small)
            y += 52
    else:
        draw.text((24, 150), "Reports", fill="#0d1f2d", font=bold)
        draw.rounded_rectangle((20, 205, 278, 260), radius=5, fill="white", outline="#c9d8d2")
        draw.text((36, 222), "Report type: Combined Report", fill="#33403c", font=small)
        draw.rounded_rectangle((20, 286, 278, 341), radius=5, fill="white", outline="#c9d8d2")
        draw.text((36, 303), "Circle: All circles", fill="#33403c", font=small)
        draw.text((24, 378), "Dates to score", fill="#557b70", font=small)
        for i, val in enumerate(("2026-08-19", "2026-08-20", "2026-08-21")):
            y = 410 + i * 40
            draw.rectangle((25, y, 42, y + 17), outline="#0d9488", width=2)
            draw.text((55, y - 2), val, fill="#33403c", font=small)
        draw.rounded_rectangle((20, 615, 278, 667), radius=6, fill="#0d9488")
        draw.text((61, 632), "Generate report (.xlsx)", fill="white", font=small)
        draw.text((330, 165), "Fleet Reports", fill="#0d1f2d", font=bold)
        draw.text((330, 225), "Choose a report type on the left:", fill="#33403c", font=regular)
        bullets = ["Fleet Registry Report - live registry and exceptions", "TM Performance Report - deployment and trips", "FM Performance Report - manager scoring", "Combined Report - all operational sheets"]
        y = 285
        for text in bullets:
            draw.ellipse((335, y + 3, 347, y + 15), fill="#158b5d")
            draw.text((365, y - 2), text, fill="#33403c", font=regular)
            y += 62
        draw.rounded_rectangle((330, 565, 1200, 680), radius=8, fill="#e1f3eb", outline="#c7e4d5")
        draw.text((360, 595), "Reports are written to the reports folder.", fill="#0c5d43", font=bold)
        draw.text((360, 635), "Wait for the completion message before opening the workbook.", fill="#557b70", font=regular)
    image.save(path)


def build_manual():
    map_image = ASSET_DIR / "map_analytics.png"
    if not map_image.exists():
        map_image = ROOT / "manual_map_analytics.png"
    draw_ui_screenshot(ASSET_DIR / "fleet_table.png", "Fleet Table")
    draw_ui_screenshot(ASSET_DIR / "reports.png", "Reports")

    doc = Document()
    setup_document(doc, "LWMC Fleet Dashboard User Manual")
    add_title_page(doc)

    doc.add_heading("Contents", level=1)
    add_table(doc, ["Section", "Use it for"], [
        ("1. Purpose and scope", "Understand what the dashboard includes and what a snapshot means."),
        ("2. Daily operating procedure", "Prepare files, launch the app, validate today's data."),
        ("3. Data sources", "Import or reset source files and understand auto-detection."),
        ("4. Map & Analytics", "Read KPIs, map state, alerts, filters and exports."),
        ("5. Fleet Table", "Investigate vehicles using status and responsibility filters."),
        ("6. Reports", "Generate registry, TM, FM and combined workbooks."),
        ("7. Data definitions", "Interpret deployment, registry, working-hours and assignment fields."),
        ("8. Troubleshooting", "Resolve stale data, missing files and report errors."),
    ])
    doc.add_page_break()

    doc.add_heading("1. Purpose And Scope", level=1)
    doc.add_paragraph("LWMC Fleet Dashboard is a Lahore-focused desktop application for reviewing a live VTMS snapshot together with the fleet establishment registry, vehicle-status data, Lahore UC and zone boundaries, employee coverage, trip reports and optional GPS history.")
    add_callout(doc, "Important", "The dashboard shows a point-in-time operational snapshot. It does not by itself prove how long a vehicle remained inside or outside its assigned area. Compare dated master files when a time series is required.", AMBER)
    doc.add_heading("What the application provides", level=2)
    add_bullets(doc, [
        "Map & Analytics: KPIs, operational cards, vehicle map, analytics views, alerts and an interactive fleet table.",
        "Fleet Table: every registered vehicle joined to the current live VTMS snapshot, with searchable filters and a fleet registry report button.",
        "Reports: Fleet Registry, TM Performance, FM Performance and Combined workbooks saved in the reports folder.",
        "Persistent data source controls: choose files from the interface, or use automatic file detection when the source names follow the documented patterns.",
    ])

    doc.add_heading("2. Daily Operating Procedure", level=1)
    doc.add_paragraph("Use this sequence at the start of each operating day. It keeps the data date, the displayed snapshot and the generated reports aligned.")
    add_numbered(doc, [
        "Download the current VTMS active-log CSV and vehicle-status export. Add the employee roster, fleet registry, trip reports or vehicle-history CSV when those sources are available.",
        "Place the files in the project folder or select them using the Import buttons at the top of the application. Keep old exports in a separate archive folder if automatic detection is being used.",
        "Launch LWMC_Fleet_Dashboard.exe from the packaged folder. Keep the executable and its _internal folder together.",
        "Wait for loading to finish. The app regenerates derived reference data and the dashboard snapshot in the background.",
        "Check the snapshot date, source labels, vehicle count and the Map & Analytics view before making an operational decision.",
        "Use filters or reports for the question at hand. Allow report generation to finish before opening the workbook or closing the application.",
    ])
    add_callout(doc, "Expected behavior", "The app may write dashboard_data.json, the standalone HTML dashboard and reference data next to the executable. This is normal and indicates that the current source set was processed.")
    doc.add_heading("Launch checks", level=2)
    add_table(doc, ["Check", "Expected result", "What to do if it fails"], [
        ("Snapshot date", "Matches today's VTMS export date.", "Review the selected VTMS file and its filename/date."),
        ("Source strip", "Each required source shows a filename or auto-detected status.", "Use the matching Import button."),
        ("Vehicle count", "A plausible count for the Lahore fleet.", "Check whether the source was filtered or stale."),
        ("Map & Analytics", "Overview cards and fleet rows are visible.", "Confirm dashboard_data.json and HTML exist; reload sources."),
    ])

    doc.add_heading("3. Data Sources", level=1)
    doc.add_paragraph("The source strip is visible above all three tabs. Selecting a file is remembered for future launches and automatically reloads the application.")
    add_table(doc, ["Button", "Accepted source", "Used for"], [
        ("Import VTMS Export", "CSV active logs", "Live position, status, UC and zone."),
        ("Import Vehicle Status", "XLS/XLSX export", "Vehicle-status and battery fields."),
        ("Import Employee Roster", "XLSX roster", "TM, FM, ZO, AM Yard and MVI responsibility."),
        ("Import Trip Reports", "One or more XLSX files", "TM/FM trip scoring by date and circle."),
        ("Import Fleet Registry", "XLSX establishment registry", "Registered category, town and fleet targets."),
        ("Import Vehicle Tracks", "CSV GPS history", "Historical route tracks and activity context."),
    ])
    doc.add_heading("Automatic detection", level=2)
    doc.add_paragraph("When no override is selected, the application looks for recognized filename patterns and generally uses the most recently modified matching file. Use Reset to auto-detect after a one-off import, or when the status strip points to the wrong file.")
    add_callout(doc, "Good practice", "Do not leave multiple competing current-day exports in the same folder when relying on auto-detection. A clean input folder makes the selected source obvious.")

    doc.add_heading("4. Map & Analytics", level=1)
    doc.add_paragraph("Map & Analytics is the operational overview. It combines the current snapshot with summary metrics, map layers, analytics cards, alerts and a vehicle table. The standalone HTML dashboard can also be opened in a browser; map tiles require internet access, while the embedded data and tables work offline.")
    add_figure(doc, map_image, "Figure 1. Current Map & Analytics view captured from the dashboard.", 6.85)
    doc.add_heading("How to read the overview", level=2)
    add_bullets(doc, [
        "Snapshot: confirms the data date currently displayed.",
        "Full fleet / Roster-scoped: Full fleet includes Lahore vehicles; Roster-scoped limits the view to vehicles inside UCs covered by the employee roster.",
        "Moving and Still / parked legend: explains the primary live-status colors.",
        "Fleet health: a composite score based on moving, online and battery-OK proportions for vehicles currently in view.",
        "KPI cards and table: provide counts and a fast way to open the written vehicle record below the dashboard.",
    ])
    doc.add_heading("Navigation and investigation", level=2)
    add_numbered(doc, [
        "Use Overview for the fleet-wide signal and Analytics for comparisons such as town targets, working-hours distribution and top zones by distance.",
        "Use Map to inspect vehicle locations, boundaries and route history. Select a vehicle or use the search field to focus the view.",
        "Use Reports for the written export. Use Fleet table when the question needs a precise combination of status, category, town, zone, UC or employee.",
        "Use the Full fleet / Roster-scoped toggle when you need either the organization-wide picture or only the staff-covered operational area.",
    ])
    doc.add_heading("Roster-scoped mode", level=2)
    doc.add_paragraph("Roster-scoped mode restricts map, KPIs, charts, table, written report and Excel export to vehicles in UCs that are covered by the roster. The Employee selector drills into one person's resolved UCs. An unresolved employee remains visible and is labeled as unresolved area so coverage gaps can be corrected in the source roster.")

    doc.add_heading("5. Fleet Table", level=1)
    doc.add_paragraph("Use Fleet Table for vehicle-level work: finding a moving vehicle, checking a town or zone, reviewing registry-only vehicles, or preparing the fleet registry report.")
    add_figure(doc, ASSET_DIR / "fleet_table.png", "Figure 2. Fleet Table with the live filter sidebar and vehicle rows.", 6.85)
    doc.add_heading("Apply filters", level=2)
    add_numbered(doc, [
        "Choose a Status: All, Moving, Still or Not Live (registry only).",
        "Choose a Vehicle Type, Town, Zone or UC. Each dropdown supports type-to-search.",
        "Choose an Employee to show vehicles currently reporting inside that person's resolved UCs. The note beneath the filters reports the person's role, coverage and moving/still counts.",
        "Read the summary after every change. It reports vehicles shown, live count, moving/still counts, registry count and registry-only count.",
        "Use Reset filters to return every selector to All.",
    ])
    add_callout(doc, "Filter logic", "Town, zone, UC and employee filters compose with AND logic. For example, Town + Zone shows only rows satisfying both selections.")
    doc.add_heading("Generate a Fleet Registry report", level=2)
    doc.add_paragraph("The button at the bottom of the Fleet Table sidebar generates the current Fleet Registry workbook from the live snapshot. The report includes the registry/live vehicle table, town summary, employee coverage and status exceptions. The result is written to the reports folder and a completion dialog confirms the path.")

    doc.add_heading("6. Reports", level=1)
    doc.add_paragraph("The Reports tab creates Excel workbooks in the reports folder. Select the report type on the left, optionally set dates or a circle for performance reports, then click Generate report (.xlsx).")
    add_figure(doc, ASSET_DIR / "reports.png", "Figure 3. Reports tab with report type, date and circle controls.", 6.85)
    add_table(doc, ["Report", "When to use it", "Main contents"], [
        ("Fleet Registry", "Current live operational snapshot.", "Every vehicle, town summary, employee coverage and exceptions."),
        ("TM Performance", "Score Town Managers.", "Deployment marks and trip-weightage marks by day."),
        ("FM Performance", "Score Fleet Managers.", "Same scoring approach for Fleet Managers."),
        ("Combined", "Share one complete workbook.", "Overview, registry, TM performance and FM performance sheets."),
    ])
    doc.add_heading("Date and circle filters", level=2)
    add_bullets(doc, [
        "Date and circle controls appear for TM, FM and Combined reports. They are hidden for Fleet Registry because that report always uses the current live snapshot.",
        "Check one or more dates to score only those dates. Leave every date unchecked to include all available trip-report dates.",
        "Choose All circles or a specific circle. Date and circle filters work together.",
        "Wait for the status message and completion dialog before opening the workbook. Close an existing workbook if Windows reports that the output file is locked.",
    ])
    doc.add_heading("Where files go", level=2)
    doc.add_paragraph("Generated workbooks are written to the reports folder in the project directory. The Reports tab includes Open reports folder for quick access.")

    doc.add_heading("7. Data Definitions", level=1)
    add_table(doc, ["Term", "Meaning"], [
        ("Live", "A vehicle present in the current VTMS snapshot."),
        ("Moving", "Current VTMS status indicates movement."),
        ("Still", "Current VTMS status indicates the vehicle is not moving."),
        ("Registry only", "Registered vehicle with no current live row."),
        ("Town target", "Target fleet count from the establishment registry; loader-rickshaw and full-fleet targets are separate."),
        ("Assigned town", "Registry town when the vehicle ID matches; a coarser tehsil-code fallback may be used otherwise."),
        ("In assigned area", "Whether the current GPS point is inside the vehicle's assigned town; it is a current status, not a duration."),
        ("Working hours", "Per-vehicle working time bucketed as Below 2h, 2-4h, 4-6h, 6-8h or Above 8h."),
        ("Zone", "A fixed group of Lahore UCs. Zone filters compose with town and UC filters."),
    ])
    doc.add_heading("Performance scoring", level=2)
    doc.add_paragraph("TM and FM performance uses a score out of 100: 20 deployment marks plus 80 trip-weightage marks. The score is evaluated per day against the manager's assigned zones and UCs, using the trip-report workbooks selected for the report.")
    add_callout(doc, "Interpret carefully", "A current 'in assigned area' flag is not a historical compliance percentage. Use stacked master_<date>.csv files or dated reports when looking for trends.", AMBER)

    doc.add_heading("8. Troubleshooting", level=1)
    add_table(doc, ["Symptom", "Likely cause", "Resolution"], [
        ("Numbers look stale", "The wrong or older source file was selected.", "Check the source strip, use Import VTMS Export, then wait for reload."),
        ("Map is empty or missing", "dashboard_data.json or the dashboard HTML is missing or could not be rebuilt.", "Confirm the files are beside the executable and reselect the VTMS source."),
        ("Vehicle status fields are blank", "The vehicle-status export is missing or has a changed layout.", "Import a valid XLS/XLSX status export and reload."),
        ("Employee list is empty", "Roster is missing, unreadable or does not match the expected sheets.", "Import the employee roster and confirm its Town Managers, Fleet Managers, Zonal Officers, AM Yards and MVI sheets."),
        ("Trip dates are empty", "No trip-report workbooks were found.", "Import one or more trip-report XLSX files or place them in New folder."),
        ("Report generation fails", "Output workbook is locked or input layout is incompatible.", "Close the workbook, verify source files and retry. Review the error dialog for the exact source."),
        ("Some staff show unresolved area", "The roster's town/area text could not map to Lahore UC geometry.", "Correct the source roster's area text and rebuild; unresolved staff remain visible by design."),
    ])
    doc.add_heading("Support checklist", level=2)
    add_bullets(doc, [
        "Record the snapshot date shown in the app.",
        "Record the source filenames displayed in the source strip.",
        "Capture the exact error message or report name.",
        "Keep the relevant source files and generated report together when escalating an issue.",
    ])

    doc.add_heading("Appendix A. Recognized File Locations", level=1)
    add_table(doc, ["Location", "Contents"], [
        ("Project root", "Executable, source exports, dashboard HTML and JSON."),
        ("New folder", "Daily Rickshaw Trips Report workbooks."),
        ("Reference Documents", "town_targets.json and other cleaned reference data."),
        ("reports", "Generated Excel workbooks."),
        ("_internal", "Packaged runtime files. Do not move or edit independently of the executable."),
    ])
    doc.add_paragraph("For scripted operation, the packaged project also supports the documented build_master and report-generation commands in README.txt. The graphical workflow is recommended for day-to-day users.")
    doc.save(MANUAL)


def build_quick_reference():
    doc = Document()
    setup_document(doc, "LWMC Fleet Dashboard Quick Reference")
    p = doc.add_paragraph(style="Title")
    p.add_run("Fleet Dashboard\nQuick Reference")
    doc.add_paragraph("One-page operating aid for daily Lahore fleet review.")
    add_callout(doc, "Daily loop", "Drop today's exports -> launch the executable -> verify date and counts -> investigate -> generate the required workbook.")
    doc.add_heading("Before opening", level=1)
    add_bullets(doc, [
        "Keep the executable and _internal folder together.",
        "Place the current VTMS CSV and vehicle-status export in the project folder, or use Import at the top of the app.",
        "Add the roster, registry, trip reports and GPS history when needed.",
    ])
    doc.add_heading("Three tabs", level=1)
    add_table(doc, ["Tab", "Use it for"], [
        ("Map & Analytics", "Fleet health, map, alerts, analytics, roster scope and quick export."),
        ("Fleet Table", "Vehicle-level filters: status, type, town, zone, UC and employee."),
        ("Reports", "Fleet Registry, TM, FM or Combined Excel workbooks."),
    ])
    doc.add_heading("Fast checks", level=1)
    add_table(doc, ["Check", "Expected"], [
        ("Snapshot", "Today's VTMS date."),
        ("Source strip", "Required files show a filename or valid auto-detected source."),
        ("Fleet count", "Plausible Lahore count."),
        ("Map", "Overview cards and rows are visible."),
    ])
    doc.add_heading("Report choices", level=1)
    add_bullets(doc, [
        "Fleet Registry: current registry/live view and exceptions.",
        "TM or FM Performance: manager scoring; optionally choose dates and circle.",
        "Combined: overview plus registry, TM and FM sheets.",
        "Leave all dates unchecked to include every available date. Open output from reports folder after completion.",
    ])
    add_callout(doc, "Trouble", "Stale numbers: check the selected source file. Empty map: confirm dashboard HTML and JSON. Failed report: close any open output workbook and retry.", AMBER)
    doc.add_paragraph("Scope: Lahore only | Generated 11 September 2026")
    doc.save(QUICK)


if __name__ == "__main__":
    build_manual()
    build_quick_reference()
    print(f"Wrote {MANUAL}")
    print(f"Wrote {QUICK}")