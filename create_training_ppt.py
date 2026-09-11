from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("LWMC_Fleet_Dashboard_Training.pptx")
SW, SH = 13.333, 7.5
NAVY = RGBColor(13, 31, 45)
INK = RGBColor(27, 43, 52)
MUTED = RGBColor(92, 111, 120)
PAPER = RGBColor(248, 250, 248)
WHITE = RGBColor(255, 255, 255)
GREEN = RGBColor(21, 139, 93)
GREEN_DARK = RGBColor(12, 93, 67)
MINT = RGBColor(225, 243, 235)
AMBER = RGBColor(226, 151, 47)
RED = RGBColor(185, 58, 57)
LINE = RGBColor(218, 228, 224)
FONT = "Aptos"
FONT_BOLD = "Aptos Display"

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
blank = prs.slide_layouts[6]


def shape(slide, kind, x, y, w, h, fill, line=None, radius=False):
    shp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line or fill
    if radius and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = 0.08
    return shp


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False,
         align=PP_ALIGN.LEFT, font=FONT, margin=0.03, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def title(slide, kicker, heading, page):
    text(slide, 0.7, 0.42, 8.8, 0.25, kicker.upper(), 10, GREEN_DARK, True)
    text(slide, 0.7, 0.72, 11.8, 0.62, heading, 28, NAVY, True, font=FONT_BOLD)
    text(slide, 12.0, 0.48, 0.65, 0.3, f"{page:02d}", 11, MUTED, True, align=PP_ALIGN.RIGHT)
    shape(slide, MSO_SHAPE.RECTANGLE, 0.7, 1.42, 1.0, 0.06, GREEN)


def footer(slide, page):
    shape(slide, MSO_SHAPE.RECTANGLE, 0.7, 7.12, 11.93, 0.012, LINE)
    text(slide, 0.7, 7.18, 7.5, 0.16, "LWMC Fleet Dashboard  |  Operator guide", 8.5, MUTED)
    text(slide, 11.2, 7.18, 1.43, 0.16, f"{page} / 11", 8.5, MUTED, align=PP_ALIGN.RIGHT)


def add_bullets(slide, x, y, w, items, size=15, color=INK, gap=0.48):
    for i, item in enumerate(items):
        yy = y + i * gap
        shape(slide, MSO_SHAPE.OVAL, x, yy + 0.09, 0.1, 0.1, GREEN)
        text(slide, x + 0.22, yy, w - 0.22, 0.32, item, size, color)


def card(slide, x, y, w, h, heading, body, accent=GREEN):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, WHITE, LINE, True)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.08, h, accent)
    text(slide, x + 0.26, y + 0.22, w - 0.45, 0.34, heading, 15, NAVY, True)
    text(slide, x + 0.26, y + 0.72, w - 0.45, h - 0.88, body, 12.5, MUTED)


# 1
slide = prs.slides.add_slide(blank)
shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, NAVY)
shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.2, SH, GREEN)
text(slide, 0.85, 1.0, 5.7, 0.32, "LWMC OPERATIONS", 12, RGBColor(133, 224, 181), True)
text(slide, 0.85, 1.55, 8.8, 1.25, "Fleet Dashboard", 42, WHITE, True, font=FONT_BOLD)
text(slide, 0.9, 2.95, 7.6, 0.82, "A practical guide to the live map, fleet table and reporting workflow.", 21, RGBColor(211, 227, 220))
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.35, 3.2, 0.62, GREEN_DARK, GREEN_DARK, True)
text(slide, 1.15, 5.52, 2.7, 0.25, "Operator training | Lahore", 12, WHITE, True)
# abstract dashboard motif
for i, (x, y, w, h, c) in enumerate([(9.4, 1.1, 2.8, 1.2, GREEN_DARK), (8.6, 2.65, 3.6, 0.55, MINT), (9.25, 3.5, 2.95, 1.65, GREEN), (8.45, 5.55, 3.7, 0.3, AMBER)]):
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, c, c, True)
text(slide, 0.9, 6.82, 6.5, 0.2, "Use this deck as a handover guide for daily operations.", 10, RGBColor(165, 188, 179))

# 2
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Purpose", "One operating picture for the fleet", 2)
text(slide, 0.7, 1.78, 6.4, 0.7, "The app joins today's live VTMS snapshot with the fleet registry, Lahore boundaries, staff roster and historical tracks.", 18, INK)
card(slide, 0.7, 3.0, 3.7, 2.4, "LIVE MAP", "See moving and still vehicles, town boundaries, route history and operational KPIs in one view.", GREEN)
card(slide, 4.82, 3.0, 3.7, 2.4, "FLEET TABLE", "Filter every registered or live vehicle by status, type, town, zone, UC and assigned staff.", AMBER)
card(slide, 8.94, 3.0, 3.7, 2.4, "REPORTS", "Generate registry, TM, FM or combined workbooks into the reports folder.", GREEN_DARK)
footer(slide, 2)

# 3
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Launch", "Open the application from the packaged folder", 3)
text(slide, 0.7, 1.76, 11.7, 0.5, "Keep the executable and its _internal folder together. Launch the executable from the same folder every day.", 16, MUTED)
card(slide, 0.85, 2.7, 3.65, 2.55, "1  LOCATE", "Open the folder containing LWMC_Fleet_Dashboard.exe and _internal.", GREEN)
card(slide, 4.85, 2.7, 3.65, 2.55, "2  LAUNCH", "Double-click the executable. The app opens with Map & Analytics selected.", GREEN)
card(slide, 8.85, 2.7, 3.65, 2.55, "3  CONFIRM", "Check the source status strip and snapshot date before using the dashboard.", AMBER)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.85, 5.75, 11.65, 0.7, MINT, MINT, True)
text(slide, 1.12, 5.96, 11.1, 0.25, "Expected on first launch: a few reference files may be written next to the executable. This is normal.", 13, GREEN_DARK, True)
footer(slide, 3)

# 4
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Start here", "The app is already configured", 4)
add_bullets(slide, 0.9, 1.95, 6.0, ["The Lahore map, fleet registry and zone data are built into the app.", "There is no setup wizard on a fresh copy.", "The app detects the most recently modified matching source file.", "Reference files are created beside the executable, never inside _internal."], 16, INK, 0.76)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 8.0, 2.0, 4.2, 3.5, NAVY, NAVY, True)
text(slide, 8.4, 2.42, 3.3, 0.3, "CHECK BEFORE WORK", 11, RGBColor(133, 224, 181), True)
text(slide, 8.4, 2.95, 3.2, 1.4, "Snapshot date\nSource status\nVehicle count\nMap loaded", 21, WHITE, True)
text(slide, 8.4, 4.72, 3.2, 0.38, "These four signals tell you whether today's data is active.", 11.5, RGBColor(211, 227, 220))
footer(slide, 4)

# 5
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Daily routine", "Drop today's exports, then launch", 5)
text(slide, 0.7, 1.72, 11.8, 0.48, "Place live source files next to the executable before opening the app. Filename patterns drive auto-detection.", 15, MUTED)
headers = ["Data source", "Recognized filename", "Used for"]
rows = [
    ("VTMS active logs", "export-vtms-active-logs-*.csv", "Live position and movement"),
    ("Vehicle status", "vs.xlsx or Vehicle_Status_*.xlsx", "Battery and reporting fields"),
    ("Employee roster", "employee*data*.xlsx", "TM / FM / ZO coverage"),
    ("Vehicle history", "activity_*.csv", "Recent route tracks"),
    ("Trip reports", "New folder/*.xlsx", "TM / FM performance"),
]
tx, ty, widths = 0.7, 2.55, [3.0, 4.4, 4.53]
for i, head in enumerate(headers):
    x = tx + sum(widths[:i]); shape(slide, MSO_SHAPE.RECTANGLE, x, ty, widths[i], 0.55, NAVY, NAVY)
    text(slide, x + 0.16, ty + 0.14, widths[i] - 0.3, 0.25, head, 11.5, WHITE, True)
for r, row in enumerate(rows):
    y = ty + 0.55 + r * 0.63
    for i, value in enumerate(row):
        x = tx + sum(widths[:i]); fill = WHITE if r % 2 == 0 else RGBColor(239, 245, 242)
        shape(slide, MSO_SHAPE.RECTANGLE, x, y, widths[i], 0.63, fill, LINE)
        text(slide, x + 0.16, y + 0.18, widths[i] - 0.3, 0.25, value, 11, INK if i == 0 else MUTED, i == 0)
text(slide, 0.7, 6.45, 11.8, 0.3, "The most recently modified matching file is selected. Keep old exports in a separate archive folder to avoid ambiguity.", 11.5, MUTED)
footer(slide, 5)

# 6
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Workflow", "From export to an operational decision", 6)
steps = [("1", "Export", "Download today's files"), ("2", "Drop", "Place them beside the exe"), ("3", "Launch", "Open the dashboard"), ("4", "Verify", "Check date and counts"), ("5", "Use", "Filter, inspect, report")]
for i, (number, head, body) in enumerate(steps):
    x = 0.72 + i * 2.48
    if i < len(steps) - 1:
        shape(slide, MSO_SHAPE.RECTANGLE, x + 1.15, 3.05, 1.35, 0.04, LINE, LINE)
    shape(slide, MSO_SHAPE.OVAL, x, 2.55, 0.9, 0.9, GREEN if i < 4 else AMBER, GREEN if i < 4 else AMBER)
    text(slide, x, 2.77, 0.9, 0.25, number, 18, WHITE, True, align=PP_ALIGN.CENTER)
    text(slide, x - 0.28, 3.75, 1.45, 0.3, head, 15, NAVY, True, align=PP_ALIGN.CENTER)
    text(slide, x - 0.42, 4.25, 1.72, 0.6, body, 11.5, MUTED, align=PP_ALIGN.CENTER)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.1, 5.65, 11.0, 0.7, MINT, MINT, True)
text(slide, 1.38, 5.88, 10.5, 0.24, "The app rebuilds the working snapshot at launch. You do not need to import a file from inside the interface.", 13, GREEN_DARK, True, align=PP_ALIGN.CENTER)
footer(slide, 6)

# 7
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Tab 1 of 3", "Map & Analytics", 7)
card(slide, 0.8, 1.95, 3.7, 3.75, "READ THE LIVE PICTURE", "Every vehicle is plotted with moving and still states. Use the map to see where the fleet is concentrated and where it is outside its assigned town.", GREEN)
card(slide, 4.82, 1.95, 3.7, 3.75, "USE THE CONTEXT", "Town and zone boundaries, KPIs and historical tracks help turn a point on the map into an operational explanation.", AMBER)
card(slide, 8.84, 1.95, 3.7, 3.75, "ASK BETTER QUESTIONS", "Which vehicles are moving? Which town is under-deployed? Is a vehicle working outside its home area?", GREEN_DARK)
footer(slide, 7)

# 8
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Tab 2 of 3", "Fleet Table", 8)
text(slide, 0.7, 1.75, 11.8, 0.5, "Use the table when the question is about a specific vehicle, assignment or exception.", 16, MUTED)
add_bullets(slide, 0.9, 2.65, 5.7, ["Filter by status, category and town.", "Narrow by zone, UC or responsible employee.", "Compare live VTMS rows with registered fleet rows.", "Use the summary strip to see how filters change the population."], 16, INK, 0.68)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 7.2, 2.35, 5.0, 2.95, NAVY, NAVY, True)
text(slide, 7.55, 2.7, 4.25, 0.25, "FILTER EXAMPLE", 11, RGBColor(133, 224, 181), True)
text(slide, 7.55, 3.2, 4.2, 1.1, "Status: moving\nTown: Wagha Town\nCategory: Loader Rickshaw", 18, WHITE, True)
text(slide, 7.55, 4.7, 4.15, 0.35, "A focused view for a quick field call.", 11.5, RGBColor(211, 227, 220))
footer(slide, 8)

# 9
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Tab 3 of 3", "Reports", 9)
card(slide, 0.8, 1.9, 2.75, 3.6, "FLEET REGISTRY", "Registry plus live status, town summary, employee coverage and exceptions.", GREEN)
card(slide, 3.82, 1.9, 2.75, 3.6, "TM PERFORMANCE", "Deployment and trip counts with date and circle filters.", AMBER)
card(slide, 6.84, 1.9, 2.75, 3.6, "FM PERFORMANCE", "Fleet Manager performance view using the same snapshot.", GREEN_DARK)
card(slide, 9.86, 1.9, 2.75, 3.6, "COMBINED", "One workbook containing overview, registry, TM and FM sheets.", NAVY)
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 5.95, 11.8, 0.62, MINT, MINT, True)
text(slide, 1.05, 6.15, 11.3, 0.22, "Reports are written to the reports folder and generated in the background. Wait for the completion message before opening the workbook.", 12.5, GREEN_DARK, True, align=PP_ALIGN.CENTER)
footer(slide, 9)

# 10
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, PAPER); title(slide, "Good habits", "Quick checks when something looks wrong", 10)
card(slide, 0.8, 1.85, 5.65, 4.4, "GOOD HABITS", "Drop only the day's files beside the executable.\n\nKeep historical exports in an archive folder.\n\nConfirm the snapshot date and source status before acting.\n\nLet report generation finish before closing the app.", GREEN)
card(slide, 6.88, 1.85, 5.65, 4.4, "TROUBLESHOOTING", "Numbers look stale: check which file was most recently modified.\n\nVehicle status is blank: confirm vs.xlsx or Vehicle_Status_*.xlsx exists.\n\nMap is empty: confirm the dashboard HTML and JSON are beside the executable.\n\nA report fails: close any workbook that is already open or locked.", RED)
footer(slide, 10)

# 11
slide = prs.slides.add_slide(blank); shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH, NAVY)
shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 0.2, SH, GREEN)
text(slide, 0.9, 1.15, 4.5, 0.25, "THE DAILY LOOP", 11, RGBColor(133, 224, 181), True)
text(slide, 0.9, 1.72, 10.5, 1.0, "Drop the files.\nLaunch the picture.\nAct on the signal.", 31, WHITE, True, font=FONT_BOLD)
text(slide, 0.9, 4.55, 8.8, 0.5, "Map, table and reports are three views of the same operational snapshot.", 17, RGBColor(211, 227, 220))
shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.9, 5.75, 2.0, 0.6, GREEN_DARK, GREEN_DARK, True)
text(slide, 1.1, 5.94, 1.6, 0.22, "Questions?", 12, WHITE, True)
text(slide, 9.1, 6.42, 3.3, 0.25, "LWMC Fleet Dashboard", 12, RGBColor(165, 188, 179), align=PP_ALIGN.RIGHT)

prs.save(OUT)
print(f"Wrote {OUT}")
